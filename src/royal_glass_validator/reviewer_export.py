"""One-way Excel reviewer export for Phase 4 classification history."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
import json
from pathlib import Path
from typing import Mapping
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


@dataclass(frozen=True)
class ReviewerEvidence:
    """Compact retained page evidence shown to a reviewer."""

    url: str
    title: str | None
    snippet: str


@dataclass(frozen=True)
class ReviewerExportRow:
    """All review-visible facts for one latest Source Record decision."""

    source_record_id: UUID
    classification_decision_id: UUID
    validation_run_id: UUID
    input_row_number: int
    original_values: Mapping[str, object]
    proposed_classification: str | None
    proposed_outcome: str
    confidence_band: str
    rationale: str
    tags: tuple[str, ...]
    evidence: tuple[ReviewerEvidence, ...]
    fetch_failures: tuple[str, ...]
    duplicate_link: str | None
    active_override_classification: str | None
    active_override_id: UUID | None
    active_override_rationale: str | None
    active_override_reviewer: str | None
    override_conflict: bool


_HEADERS = (
    "Source Record ID", "Classification Decision ID", "Validation Run ID", "Input Row", "Original Record", "Proposed Classification",
    "Proposed Outcome", "Confidence", "Reasons", "Tags", "Evidence", "Conflicts / Revalidation",
    "Fetch Errors", "Duplicate Link", "Protected Override", "Active Override ID", "Override Rationale", "Override Reviewer", "Reviewer Note",
)


def protected_override_conflict(proposed_classification: str | None, active_override_classification: str | None) -> bool:
    """Return whether a new automated proposal needs reviewer revalidation.

    This never modifies the protected override; it only makes disagreement visible.
    """
    return bool(proposed_classification and active_override_classification and proposed_classification != active_override_classification)


def export_reviewer_workbook(
    rows: tuple[ReviewerExportRow, ...], output_directory: Path, *, generated_at: datetime | None = None
) -> Path:
    """Create a dated, one-way reviewer workbook without changing stored decisions."""
    timestamp = generated_at or datetime.now(timezone.utc)
    output_directory.mkdir(parents=True, exist_ok=True)
    output_path = _available_output_path(output_directory, timestamp)
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Reviewer Export"
    worksheet.append(_HEADERS)
    for row in rows:
        worksheet.append(tuple(_safe_cell_value(value) for value in (
            str(row.source_record_id), str(row.classification_decision_id), str(row.validation_run_id), row.input_row_number,
            _format_original_values(row.original_values), row.proposed_classification or "", row.proposed_outcome,
            row.confidence_band, row.rationale, ", ".join(row.tags), _format_evidence(row.evidence),
            "Conflicts with protected override — revalidation required." if row.override_conflict else "",
            "\n".join(row.fetch_failures), row.duplicate_link or "", row.active_override_classification or "",
            str(row.active_override_id) if row.active_override_id else "", row.active_override_rationale or "", row.active_override_reviewer or "", "",
        )))

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:S{max(1, worksheet.max_row)}"
    worksheet.row_dimensions[1].height = 30
    for column, width in {"A": 38, "B": 38, "C": 38, "D": 12, "E": 45, "F": 22, "G": 20, "H": 14, "I": 42, "J": 24, "K": 55, "L": 38, "M": 38, "N": 38, "O": 22, "P": 38, "Q": 42, "R": 24, "S": 36}.items():
        worksheet.column_dimensions[column].width = width
    workbook.save(output_path)
    return output_path


def _available_output_path(output_directory: Path, timestamp: datetime) -> Path:
    stem = f"royal-glass-reviewer-export-{timestamp.date().isoformat()}"
    candidate = output_directory / f"{stem}.xlsx"
    sequence = 2
    while candidate.exists():
        candidate = output_directory / f"{stem}-{sequence}.xlsx"
        sequence += 1
    return candidate


def _format_original_values(values: Mapping[str, object]) -> str:
    return "\n".join(f"{key}: {_format_value(value)}" for key, value in sorted(values.items(), key=lambda item: str(item[0]).casefold()))


def _format_value(value: object) -> str:
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _format_evidence(evidence: tuple[ReviewerEvidence, ...]) -> str:
    return "\n\n".join(f"{item.title or 'Untitled'} — {item.url}\n{item.snippet}" for item in evidence)


def _safe_cell_value(value: object) -> object:
    """Keep retained reviewer content as literal cells, never executable Excel formulas."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return f"'{value}"
    return value
