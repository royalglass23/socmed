"""Validate and persist immutable Needs Validation workbook source records."""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass
from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
from typing import Any, Protocol
from uuid import UUID, uuid4


class ImportValidationError(ValueError):
    """Raised when an input workbook is missing, invalid, or ambiguous."""


@dataclass(frozen=True)
class ImportedRun:
    id: UUID
    input_workbook_name: str
    input_workbook_sha256: str
    input_provenance: Mapping[str, object]
    rule_version: str


@dataclass(frozen=True)
class ImportedSourceRecord:
    id: UUID
    validation_run_id: UUID
    entity_id: str
    source_system: str
    source_locator: str
    input_row_number: int
    original_values: Mapping[str, object]
    original_value_hash: str


@dataclass(frozen=True)
class WorkbookImportSummary:
    validation_run_id: UUID
    input_workbook_name: str
    imported_record_count: int


class ImportRepository(Protocol):
    """The atomic persistence boundary for one fully validated workbook import."""

    def import_records(self, run: ImportedRun, records: tuple[ImportedSourceRecord, ...]) -> None: ...


def import_workbook(
    workbook_path: Path,
    repository: ImportRepository,
    *,
    rule_version: str,
) -> WorkbookImportSummary:
    """Validate a workbook completely, then persist its source records once."""
    run, records = prepare_workbook_import(workbook_path, rule_version=rule_version)
    repository.import_records(run, records)
    return WorkbookImportSummary(
        validation_run_id=run.id,
        input_workbook_name=run.input_workbook_name,
        imported_record_count=len(records),
    )


def prepare_workbook_import(
    workbook_path: Path,
    *,
    rule_version: str,
) -> tuple[ImportedRun, tuple[ImportedSourceRecord, ...]]:
    """Return a validated immutable import plan without touching persistence."""
    if not rule_version.strip():
        raise ImportValidationError("A non-empty ruleset version is required for an import.")
    if not workbook_path.is_file():
        raise ImportValidationError(f"Input workbook does not exist: {workbook_path.name}")
    if workbook_path.suffix.casefold() != ".xlsx":
        raise ImportValidationError("Input workbook must be an .xlsx file.")
    try:
        workbook_bytes = workbook_path.read_bytes()
    except OSError as error:
        raise ImportValidationError(f"Could not read the input workbook: {workbook_path.name}") from error
    workbook_sha256 = sha256(workbook_bytes).hexdigest()

    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as error:
        raise ImportValidationError("Install workbook import dependencies before importing an .xlsx file.") from error

    try:
        workbook = load_workbook(BytesIO(workbook_bytes), read_only=True, data_only=False)
    except (InvalidFileException, OSError, ValueError) as error:
        raise ImportValidationError(f"Could not read the input workbook: {workbook_path.name}") from error

    try:
        worksheet = _select_needs_validation_worksheet(workbook)
        headers = _read_headers(worksheet)
        records = _read_source_records(worksheet, headers, workbook_path)
        worksheet_title = worksheet.title
    finally:
        workbook.close()

    if not records:
        raise ImportValidationError("Needs Validation worksheet contains no source records.")

    run_id = uuid4()
    run = ImportedRun(
        id=run_id,
        input_workbook_name=workbook_path.name,
        input_workbook_sha256=workbook_sha256,
        input_provenance={
            "source_system": "royal_glass_competitor_workbook",
            "worksheet": worksheet_title,
            "header_names": headers,
            "cohort": "Needs Validation",
        },
        rule_version=rule_version,
    )
    return (
        run,
        tuple(
            ImportedSourceRecord(
                id=uuid4(),
                validation_run_id=run_id,
                entity_id=entity_id,
                source_system="royal_glass_competitor_workbook",
                source_locator=worksheet_title,
                input_row_number=row_number,
                original_values=values,
                original_value_hash=_stable_value_hash(values),
            )
            for entity_id, row_number, values in records
        ),
    )


def _select_needs_validation_worksheet(workbook: Any) -> Any:
    matches = [worksheet for worksheet in workbook.worksheets if worksheet.title.strip().casefold() == "needs validation"]
    if len(matches) != 1:
        raise ImportValidationError("Input workbook must contain exactly one Needs Validation worksheet.")
    return matches[0]


def _read_headers(worksheet: Any) -> tuple[str, ...]:
    first_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    if not first_row:
        raise ImportValidationError("Needs Validation worksheet is missing its header row.")
    headers: list[str] = []
    for column_number, value in enumerate(first_row, start=1):
        if not isinstance(value, str) or not value.strip():
            raise ImportValidationError(f"Needs Validation header at column {column_number} is blank or invalid.")
        headers.append(value)
    normalized_headers = [header.strip().casefold() for header in headers]
    if len(normalized_headers) != len(set(normalized_headers)):
        raise ImportValidationError("Needs Validation worksheet has ambiguous duplicate headers.")
    required_headers = {"entity id", "decision group"}
    if not required_headers.issubset(normalized_headers):
        raise ImportValidationError("Needs Validation worksheet requires Entity ID and Decision Group columns.")
    return tuple(headers)


def _read_source_records(
    worksheet: Any,
    headers: tuple[str, ...],
    workbook_path: Path,
) -> list[tuple[str, int, Mapping[str, object]]]:
    entity_id_header = next(header for header in headers if header.strip().casefold() == "entity id")
    decision_group_header = next(header for header in headers if header.strip().casefold() == "decision group")
    records: list[tuple[str, int, Mapping[str, object]]] = []
    seen_entity_ids: set[str] = set()
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        values = {header: _json_value(value) for header, value in zip(headers, row, strict=True)}
        entity_id_value = values[entity_id_header]
        if not isinstance(entity_id_value, str) or not entity_id_value.strip():
            raise ImportValidationError(f"Needs Validation row {row_number} has no Entity ID.")
        entity_id = entity_id_value.strip()
        if entity_id in seen_entity_ids:
            raise ImportValidationError(f"Needs Validation worksheet has duplicate Entity ID: {entity_id}")
        decision_group = values[decision_group_header]
        if not isinstance(decision_group, str) or not decision_group.strip().casefold().startswith("needs validation"):
            raise ImportValidationError(
                f"Needs Validation row {row_number} is not in the Needs Validation cohort: {workbook_path.name}"
            )
        seen_entity_ids.add(entity_id)
        records.append((entity_id, row_number, values))
    return records


def _json_value(value: object) -> object:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    raise ImportValidationError(f"Workbook has an unsupported cell value type: {type(value).__name__}")


def _stable_value_hash(values: Mapping[str, object]) -> str:
    encoded = json.dumps(values, sort_keys=True, separators=(",", ":"), ensure_ascii=False, allow_nan=False).encode("utf-8")
    return sha256(encoded).hexdigest()
