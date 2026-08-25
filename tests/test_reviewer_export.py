from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from uuid import uuid4


class ReviewerWorkbookTests(unittest.TestCase):
    def test_exports_original_record_audit_evidence_and_protected_override_fields(self) -> None:
        from royal_glass_validator.reviewer_export import ReviewerEvidence, ReviewerExportRow, export_reviewer_workbook

        row = ReviewerExportRow(
            source_record_id=uuid4(),
            classification_decision_id=uuid4(),
            validation_run_id=uuid4(),
            input_row_number=12,
            original_values={"Canonical Name": "Clear View Glass", "Website": "https://clearview.example"},
            proposed_classification="direct",
            proposed_outcome="direct",
            confidence_band="high",
            rationale="Official site confirms pool-fencing installation in Auckland.",
            tags=("pool-fencing", "residential"),
            evidence=(
                ReviewerEvidence(
                    url="https://clearview.example/pool-fencing",
                    title="Pool fencing",
                    snippet="We install frameless glass pool fencing throughout Auckland.",
                ),
            ),
            fetch_failures=("timeout: https://clearview.example/contact (attempt 1)",),
            duplicate_link="High-certainty link: Clear View Glass (normalized_domain).",
            active_override_classification="adjacent",
            active_override_id=uuid4(),
            active_override_rationale="Reviewer confirmed this is a landscaping partnership, not an installer.",
            active_override_reviewer="Aroha Reviewer",
            override_conflict=True,
        )

        with TemporaryDirectory() as temporary_directory:
            output_path = export_reviewer_workbook((row,), Path(temporary_directory), generated_at=datetime(2026, 8, 25, 1, 2, tzinfo=timezone.utc))

            from openpyxl import load_workbook

            workbook = load_workbook(output_path)
            worksheet = workbook["Reviewer Export"]
            headers = [cell.value for cell in worksheet[1]]
            values = [cell.value for cell in worksheet[2]]

        self.assertEqual(output_path.name, "royal-glass-reviewer-export-2026-08-25.xlsx")
        self.assertEqual(headers, [
            "Source Record ID",
            "Classification Decision ID",
            "Validation Run ID",
            "Input Row",
            "Original Record",
            "Proposed Classification",
            "Proposed Outcome",
            "Confidence",
            "Reasons",
            "Tags",
            "Evidence",
            "Conflicts / Revalidation",
            "Fetch Errors",
            "Duplicate Link",
            "Protected Override",
            "Active Override ID",
            "Override Rationale",
            "Override Reviewer",
            "Reviewer Note",
        ])
        self.assertIn("Canonical Name: Clear View Glass", values[4])
        self.assertEqual(values[5:10], ["direct", "direct", "high", "Official site confirms pool-fencing installation in Auckland.", "pool-fencing, residential"])
        self.assertIn("Pool fencing", values[10])
        self.assertEqual(values[11], "Conflicts with protected override — revalidation required.")
        self.assertIn("timeout:", values[12])
        self.assertEqual(values[14], "adjacent")
        self.assertEqual(values[15], str(row.active_override_id))
        self.assertIsNone(values[18])

    def test_writes_untrusted_reviewer_content_as_literal_cells_not_excel_formulas(self) -> None:
        from royal_glass_validator.reviewer_export import ReviewerEvidence, ReviewerExportRow, export_reviewer_workbook

        row = ReviewerExportRow(
            source_record_id=uuid4(), classification_decision_id=uuid4(), validation_run_id=uuid4(), input_row_number=1,
            original_values={"=Untrusted field name": "value"}, proposed_classification="direct",
            proposed_outcome="direct", confidence_band="high", rationale="@untrusted rationale", tags=("+tag",),
            evidence=(ReviewerEvidence(url="https://example.test", title="=Untrusted title", snippet="literal evidence"),),
            fetch_failures=(), duplicate_link=None, active_override_classification=None,
            active_override_id=None, active_override_rationale=None, active_override_reviewer=None, override_conflict=False,
        )

        with TemporaryDirectory() as temporary_directory:
            output_path = export_reviewer_workbook((row,), Path(temporary_directory))
            from openpyxl import load_workbook

            worksheet = load_workbook(output_path).active
            self.assertEqual(worksheet["E2"].data_type, "s")
            self.assertTrue(worksheet["E2"].value.startswith("'=Untrusted field name:"))
            self.assertEqual(worksheet["I2"].value, "'@untrusted rationale")
            self.assertEqual(worksheet["J2"].value, "'+tag")
            self.assertTrue(worksheet["K2"].value.startswith("'=Untrusted title"))

    def test_marks_later_conflicting_automation_for_revalidation_without_replacing_the_override(self) -> None:
        from royal_glass_validator.reviewer_export import protected_override_conflict

        self.assertTrue(protected_override_conflict("direct", "adjacent"))
        self.assertFalse(protected_override_conflict("direct", "direct"))
        self.assertFalse(protected_override_conflict("direct", None))
